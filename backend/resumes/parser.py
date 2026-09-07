"""
Reusable parser service: extracts raw text from PDF and DOCX resumes.
Kept independent of Django models so it's easy to unit test in isolation.
"""
from pypdf import PdfReader
from docx import Document


class ParsingError(Exception):
    """Raised when a resume file cannot be parsed."""
    pass


def parse_pdf(file_obj) -> str:
    try:
        reader = PdfReader(file_obj)
        if len(reader.pages) == 0:
            raise ParsingError("The PDF file appears to be empty.")

        text_chunks = []
        for page in reader.pages:
            text_chunks.append(page.extract_text() or "")

        full_text = "\n".join(text_chunks).strip()

        if not full_text:
            # Likely a scanned/image-based PDF with no extractable text layer.
            raise ParsingError(
                "No selectable text found in this PDF. It may be a scanned "
                "document. OCR support can be added in a future version."
            )
        return full_text
    except ParsingError:
        raise
    except Exception as exc:
        raise ParsingError(f"Could not parse PDF file: {exc}")


def parse_docx(file_obj) -> str:
    try:
        document = Document(file_obj)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

        # Also pull text out of tables, since resumes sometimes use table layouts
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)

        full_text = "\n".join(paragraphs).strip()
        if not full_text:
            raise ParsingError("The DOCX file appears to be empty.")
        return full_text
    except ParsingError:
        raise
    except Exception as exc:
        raise ParsingError(f"Could not parse DOCX file: {exc}")


def extract_text(file_obj, file_type: str) -> str:
    """
    Single entry point used by the rest of the app.
    file_type must be 'pdf' or 'docx'.
    """
    file_obj.seek(0)
    if file_type == 'pdf':
        return parse_pdf(file_obj)
    elif file_type == 'docx':
        return parse_docx(file_obj)
    else:
        raise ParsingError(f"Unsupported file type: {file_type}")
